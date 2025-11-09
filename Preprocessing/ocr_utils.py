import os
import fitz  # PyMuPDF
import docx
from PIL import Image
from io import BytesIO
import pytesseract
import zipfile
import xml.etree.ElementTree as ET

# مسیر اجرایی tesseract OCR را مشخص می‌کند
pytesseract.pytesseract.tesseract_cmd = "tesseract"

def extract_text_from_image_bytes(img_bytes):
    """متن‌کشی از تصویر (با استفاده از OCR)"""
    img = Image.open(BytesIO(img_bytes))
    return pytesseract.image_to_string(img, lang='fas+eng')

def extract_pdf_text_and_images(filepath: str) -> list[str]:
    """استخراج متن و تصاویر از فایل PDF و OCR روی تصاویر"""
    lines = []
    doc = fitz.open(filepath)
    for page in doc:
        # استخراج متن صفحه
        lines.extend(page.get_text("text").splitlines())
        # استخراج تصاویر و OCR روی آن‌ها
        for img in page.get_images(full=True):
            xref = img[0]
            base_image = doc.extract_image(xref)
            ocr_text = extract_text_from_image_bytes(base_image["image"])
            if ocr_text.strip():
                lines.append(f"[IMAGE_TEXT]: {ocr_text.strip()}")
    return lines

def extract_docx_text_and_images(filepath: str) -> list[str]:
    """استخراج متن، تصاویر و SmartArt از فایل DOCX"""
    lines = []

    # استخراج متن از SmartArt (اگر وجود داشته باشد)
    smartart_lines = extract_smartart_text(filepath)
    if smartart_lines:
        lines.append("[SMARTART_TEXT]:")
        lines.extend(smartart_lines)

    # استخراج متن اصلی فایل
    doc = docx.Document(filepath)
    rels = doc.part._rels

    for para in doc.paragraphs:
        if para.text.strip():
            lines.append(para.text.strip())

    # پیدا کردن تصاویر و OCR روی آن‌ها
    for para in doc.paragraphs:
        for run in para.runs:
            if 'imagedata' in run._element.xml or 'graphic' in run._element.xml:
                for rel in rels:
                    rel_obj = rels[rel]
                    if "image" in rel_obj.target_ref:
                        try:
                            ocr_text = extract_text_from_image_bytes(rel_obj.target_part.blob)
                            if ocr_text.strip():
                                lines.append(f"[IMAGE_TEXT]: {ocr_text.strip()}")
                        except Exception as e:
                            print(f"⚠️ OCR failed for image in {filepath}: {e}")
    return lines

def extract_smartart_text(docx_path: str) -> list[str]:
    """استخراج متن‌های موجود در قسمت SmartArt فایل DOCX"""
    smartart_text = []
    try:
        with zipfile.ZipFile(docx_path, 'r') as docx_zip:
            # پیدا کردن فایل‌های XML مربوط به SmartArt
            diagram_files = [name for name in docx_zip.namelist() if name.startswith("word/diagrams/data")]
            for diagram_file in diagram_files:
                xml_content = docx_zip.read(diagram_file)
                tree = ET.fromstring(xml_content)
                for elem in tree.iter():
                    if elem.tag.endswith('t') and elem.text:
                        smartart_text.append(elem.text.strip())
    except Exception as e:
        print(f"⚠️ SmartArt extraction failed in {docx_path}: {e}")
    return smartart_text

def extract_text_and_images(filepath: str) -> list[str]:
    """انتخاب روش استخراج بسته به نوع فایل"""
    ext = os.path.splitext(filepath)[-1].lower()
    if ext == '.pdf':
        return extract_pdf_text_and_images(filepath)
    elif ext == '.docx':
        return extract_docx_text_and_images(filepath)
    else:
        print(f"⚠️ Unsupported file type: {ext}")
        return []